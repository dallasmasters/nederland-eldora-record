"""Build the page's data as a spreadsheet: docs/eldora-record.xlsx plus docs/eldora-payments.csv (the payment ledger).

The .xlsx is only replaced when its cell contents change, so an unchanged workbook doesn't look like a new version
(a new version must also be uploaded to the Drive copy the page links to; see site_config.json data_drive_url).

Everything comes from the same sources as the page (analysis/eldora_due_from.csv, the constants in
eldora_spending_page.py, analysis/eldora_votes.csv) or from figures quoted on the page with their sources. Amounts
keep cents. Totals and differences are spreadsheet formulas so readers can check them; the program opening the file
calculates them.
"""
import csv
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import eldora_spending_page as page  # noqa: E402

ROOT = page.ROOT
OUT = os.path.join(ROOT, 'docs')
SITE = 'https://eldorarecord.org/'
CONTACT = 'corrections@eldorarecord.org'
DOC = 'https://nederlandco.civicweb.net/document/{}/'
AUDIT_2025 = 'https://www.nederlandco.gov/DocumentCenter/View/594/2025-Nederland-Financials-Audit'
TOWN_UPDATE = 'https://www.nederlandco.gov/m/newsflash/home/detail/21'
ESTIMATE = 2_500_000

USD = '"$"#,##0.00'
USD0 = '"$"#,##0'
DATE = 'mmm d, yyyy'
HEAD_FILL = PatternFill('solid', fgColor='DDE7EA')
FLAG_FILL = PatternFill('solid', fgColor='FBE3D2')
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')


def d(s):
    return datetime.strptime(s, '%Y-%m-%d').date()


def inv_date(s):
    return datetime.strptime(s, '%m/%d/%y').date()


def text(s):
    """Page notes are written in the page's shorthand; apply its American-date and AP-month rules."""
    return page.ap_months(page.american_dates(s)) if s else s


def sheet(wb, title, headers, rows, widths, formats=None, wrap=(), links=None, flag_col=None, first=False):
    """Write a table: bold frozen header, filters, column formats, optional hyperlink column."""
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = WRAP
    for r in rows:
        ws.append(r)
    formats = formats or {}
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            name = headers[c.column - 1]
            if name in formats:
                c.number_format = formats[name]
            c.alignment = WRAP if name in wrap else TOP
        if flag_col and row[headers.index(flag_col)].value:
            for c in row:
                c.fill = FLAG_FILL
    if links:   # {link-text column: url column}; the url column stays visible for copying
        for text_col, url_col in links.items():
            ti, ui = headers.index(text_col), headers.index(url_col)
            for row in ws.iter_rows(min_row=2):
                if row[ui].value:
                    row[ti].hyperlink = row[ui].value
                    row[ti].font = Font(color='1F5F8B', underline='single')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def contents(path):
    """Every cell value and hyperlink, sheet by sheet: what a reader sees, ignoring file timestamps."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    return [(ws.title, [[(c.value, c.hyperlink.target if c.hyperlink else None) for c in row] for row in ws.iter_rows()])
            for ws in wb]


def main():
    rows = list(csv.DictReader(open(os.path.join(ROOT, 'analysis', 'eldora_due_from.csv'))))
    for r in rows:
        r['amt'] = float(r['amount'])
        r['bal'] = -page.amt(r['account_balance_after'])
    meetings = sorted({r['warrant_meeting'] for r in rows})
    bal_at = {m: max(r['bal'] for r in rows if r['warrant_meeting'] == m) for m in meetings}
    src_at = {r['warrant_meeting']: r['source'] for r in rows}
    asof = meetings[-1]
    wb = Workbook()

    # ---------------------------------------------------------------- About
    ws = wb.active
    ws.title = 'About'
    about = [
        ("Nederland's Eldora acquisition: data", True),
        (f'From {SITE}, an independent compilation of Town of Nederland public records. '
         'Not affiliated with the Town of Nederland, POWDR, or Eldora Mountain Resort.', False),
        (f'Payments through the {text(page.fmt(d(asof)))}, warrant list.', False),
        ('', False),
        ('Sheets', True),
        ('Summary: the headline figures on the page, with how each is calculated.', False),
        ('Payments: every invoice charged to General Fund account 10-00-1150, "Due From" (the account used for Eldora acquisition costs), '
         'plus the amounts the Town\'s ledger shows were added to the account between warrant lists.', False),
        ('Off-list postings: amounts added to the account balance between warrant lists, with what the Town\'s ledger '
         'for the account says about each.', False),
        ('Running balance: the account balance after each warrant list the Board approved.', False),
        ('Payees: totals by payee, with the evidence rating used on the page.', False),
        ('Estimates vs paid: what the Board was told an engagement would cost, and what was paid.', False),
        ('Commitments: amounts committed or authorized, and their status.', False),
        ('CSAFE and funds: who owns the pooled cash and investments the Town says would pay the costs.', False),
        ('For scale: Town budgets and other recent discretionary spending.', False),
        ('Board votes: every Board action on the acquisition, from the adopted minutes.', False),
        ('', False),
        ('Notes', True),
        ('Amounts keep cents, as in the warrant lists. The page rounds to whole dollars, so a total here can differ '
         'from the page by a dollar.', False),
        ('Totals and differences are formulas, so you can check them. Your spreadsheet program calculates them when '
         'the file is opened.', False),
        ('Evidence ratings: "Named in Board documents" means a resolution, staff memo, or the purchase agreement states '
         'the firm\'s role in the Eldora acquisition. "Account coding only" means a regular Town vendor whose other work is billed '
         'to normal budget lines, so only the account links these charges to Eldora. "No other record" means the '
         'payee appears in no Board document searched except the warrant lists.', False),
        ('The Payments sheet gives the days from each invoice date to the warrant list that paid it; the typical '
         'gap is about 33 days.', False),
        ("Statements that something isn't in the records refer only to the records searched for the page; it may be "
         'documented elsewhere, including in earlier records.', False),
        ('Sources: Board packets and minutes on the Town CivicWeb portal, the Town website (FAQs, updates, audits), '
         'the Town\'s ledger for account 10-00-1150 released in its Sept. 25, 2026, response to a public records '
         f'request ({SITE}{page.CORA_LEDGER}), and state statutes. Each row links to its source.', False),
        ('An invoice sometimes reaches the account balance one list before it appears on a warrant list (for '
         'example, Widner Juran invoice 46.0322, $16,985). Those timing differences net to zero and are not counted '
         'as off-list postings.', False),
        (f'Corrections: {CONTACT}', False),
    ]
    for line, bold in about:
        ws.append([line])
        ws.cell(ws.max_row, 1).font = Font(bold=bold, size=14 if ws.max_row == 1 else 11)
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 110
    ws['A2'].hyperlink = SITE

    # ---------------------------------------------------------------- Payments
    pay_headers = ['Warrant list date', 'Payee (as listed)', 'Payee', 'Invoice', 'Invoice date', 'Amount',
                   'On warrant list', 'Days from invoice to warrant list', 'Evidence rating',
                   'Source', 'Source URL']
    pay_rows = []
    for r in sorted(rows + page.OFF_LIST, key=lambda x: (x['warrant_meeting'], 'label' in x)):
        name, tier, _ = page.PAYEES[r['vendor']]
        idate = inv_date(r['invoice_date'])
        if 'label' in r:
            first = r['first']
            note = r['label'] + (f'; first paid on the {text(page.fmt(d(first[0])))}, warrant list, p. {first[2]} '
                                 f'(https://nederlandco.civicweb.net/document/{first[1]}/)' if first else '')
            note += f'; first in the account balance on the {text(page.fmt(d(r["warrant_meeting"])))}, warrant list'
            pay_rows.append([None, r['vendor'], name, r['invoice'], idate, r['amt'], 'No', None,
                             page.EVIDENCE[tier], note, SITE + page.CORA_LEDGER])
            continue
        lag = (d(r['warrant_meeting']) - idate).days
        pay_rows.append([d(r['warrant_meeting']), r['vendor'], name, r['invoice'], idate, r['amt'], 'Yes', lag,
                         page.EVIDENCE[tier], f'Warrant list p. {r["warrant_page"]}', r['source']])
    n_pay = len(pay_rows)
    sheet(wb, 'Payments', pay_headers, pay_rows, [13, 26, 26, 14, 12, 14, 10, 12, 22, 34, 30, 48],
          formats={'Warrant list date': DATE, 'Invoice date': DATE, 'Amount': USD},
          wrap={'Source'}, links={'Source': 'Source URL'})
    with open(os.path.join(OUT, 'eldora-payments.csv'), 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(pay_headers)
        for r in pay_rows:
            w.writerow([v.isoformat() if isinstance(v, date) else ('' if v is None else v) for v in r])
    P = lambda col: f"Payments!${col}$2:${col}${n_pay + 1}"   # noqa: E731

    # ---------------------------------------------------------------- Off-list postings
    prev, gaps = None, []
    for m in meetings:
        listed = sum(r['amt'] for r in rows if r['warrant_meeting'] == m)
        jump = round(bal_at[m] - (bal_at[prev] if prev else 0) - listed, 2)
        known = [x for x in page.OFF_LIST if x['warrant_meeting'] == m]
        for x in known:   # named in the Town's ledger (public records response)
            first = x['first']
            what = x['label'] + (f'. Invoice {x["invoice"]} first paid on the {text(page.fmt(d(first[0])))}, warrant list'
                                 if first else '')
            gaps.append([d(prev), d(m), x['amt'], page.PAYEES[x['vendor']][0], what, 'Town ledger',
                         SITE + page.CORA_LEDGER])
        rest = round(jump - sum(x['amt'] for x in known), 2)
        if abs(rest) >= 0.01:
            gaps.append([d(prev), d(m), rest, '', 'Payee not in the records searched. The balance of account 10-00-1150, '
                         '"Due From," rose by more than the invoices on the warrant list.', 'Warrant list', src_at[m]])
        prev = m
    # an invoice can reach the balance one list before it appears on a list; a rise followed by an equal drop on the
    # next list is that timing difference, not a posting
    k = 0
    while k < len(gaps) - 1:
        if not gaps[k][3] and gaps[k][1] == gaps[k + 1][0] and abs(gaps[k][2] + gaps[k + 1][2]) < 0.01:
            del gaps[k:k + 2]
        else:
            k += 1
    sheet(wb, 'Off-list postings', ['After list of', 'Before list of', 'Amount', 'Attributed to', 'What is known',
                                    'Source', 'Source URL'], gaps, [13, 13, 14, 16, 70, 14, 48],
          formats={'After list of': DATE, 'Before list of': DATE, 'Amount': USD}, wrap={'What is known'},
          links={'Source': 'Source URL'})
    n_gap = len(gaps)

    # ---------------------------------------------------------------- Running balance
    rb = []
    for i, m in enumerate(meetings, 2):
        change = f'=B{i}-B{i - 1}' if i > 2 else f'=B{i}'
        listed = f'=SUMIFS({P("F")},{P("A")},A{i},{P("G")},"Yes")'
        rb.append([d(m), bal_at[m], change, listed, f'=C{i}-D{i}', 'Warrant list', src_at[m]])
    sheet(wb, 'Running balance', ['Warrant list date', 'Balance after list', 'Change since previous list',
                                  'Invoices on this list', 'Not on the list', 'Source', 'Source URL'],
          rb, [13, 16, 16, 16, 14, 14, 48],
          formats={'Warrant list date': DATE, **{k: USD for k in ('Balance after list', 'Change since previous list',
                                                                  'Invoices on this list', 'Not on the list')}},
          links={'Source': 'Source URL'})
    last_rb = len(rb) + 1

    # ---------------------------------------------------------------- Payees
    by_vendor = defaultdict(list)
    for r in rows + page.OFF_LIST:
        by_vendor[r['vendor']].append(r)
    unattributed = round(sum(g[2] for g in gaps if not g[3]), 2)
    pr = []
    order = sorted(by_vendor, key=lambda v: -sum(x['amt'] for x in by_vendor[v]))
    for i, v in enumerate(order, 2):
        name, tier, note = page.PAYEES[v]
        who = next((x for x in page.WHO if x[2] == v), None)
        if who:
            note = f'{who[3]} · {who[4]}'
        dates = [inv_date(x['invoice_date']) for x in by_vendor[v]]
        pr.append([name, v, f'=SUMIF({P("B")},B{i},{P("F")})', f'=COUNTIF({P("B")},B{i})', min(dates), max(dates),
                   page.EVIDENCE[tier], text(note)])
    i = len(pr) + 2
    if unattributed:
        pr.append(['Not on any warrant list', '', f"=SUMIF('Off-list postings'!$D$2:$D${n_gap + 1},\"\","
                                                 f"'Off-list postings'!$C$2:$C${n_gap + 1})",
                   f"=COUNTIF('Off-list postings'!$D$2:$D${n_gap + 1},\"\")", None, None,
                   'Payee not in the records searched', 'Posted to the account between warrant lists; see Off-list postings.'])
        i = len(pr) + 2
    pr.append(['Total', '', f'=SUM(C2:C{i - 1})', f'=SUM(D2:D{i - 1})', None, None, '', ''])
    ws = sheet(wb, 'Payees', ['Payee', 'Payee (as listed)', 'Total paid', 'Invoices', 'First invoice date',
                              'Last invoice date', 'Evidence rating', 'Note'],
               pr, [30, 28, 15, 9, 13, 13, 24, 80],
               formats={'Total paid': USD, 'First invoice date': DATE, 'Last invoice date': DATE}, wrap={'Note'})
    for c in ws[i]:
        c.font = Font(bold=True)
    ws.auto_filter.ref = f'A1:H{i - 1}'
    payee_total = f'Payees!C{i}'

    # ---------------------------------------------------------------- Summary
    paid_total = round(bal_at[asof], 2)
    sm = [
        ['Paid through warrant lists', f'=SUMIF({P("G")},"Yes",{P("F")})',
         'Sum of the invoices on Board-approved warrant lists (Payments, On warrant list = Yes).', 'Payments'],
        ['Posted between warrant lists', f"=SUM('Off-list postings'!C2:C{n_gap + 1})",
         'Amounts added to the account between warrant lists (Town ledger).', 'Off-list postings'],
        ['Already paid', '=B2+B3', f'Charged to the account through the {text(page.fmt(d(asof)))}, warrant list. Should '
         'equal the last running balance (next row).', 'Payments, Off-list postings'],
        ['Check: last running balance', f"='Running balance'!B{last_rb}",
         'The account balance on the latest warrant list.', 'Running balance'],
        ["Town's estimate if the sale fails", ESTIMATE, 'Transaction costs the Town says it would owe, paid from CSAFE '
         'investments (Sept. 30, 2026, update).', TOWN_UPDATE],
        ['Not yet paid (estimate minus already paid)', '=B6-B4', "Reads the Town's $2.5 million as the total. Not "
         'itemized by the Town.', ''],
        ["Alternative reading: estimate plus already paid", '=B6+B4', "If the Town's $2.5 million does not include "
         'what is already paid.', ''],
        ['Paid to POWDR (the seller)', f'=SUMIF({P("B")},"POWDR Corp",{P("F")})', 'Invoices on warrant lists.', 'Payments'],
        ['Check: total by payee', f'={payee_total}', 'Should equal Already paid.', 'Payees'],
    ]
    ws = sheet(wb, 'Summary', ['Figure', 'Amount', 'How it is calculated', 'Source'], [], [44, 16, 80, 40])
    for r in sm:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = USD
        for c in row:
            c.alignment = WRAP
        if row[3].value and row[3].value.startswith('http'):
            row[3].hyperlink = row[3].value
    ws['A4'].font = Font(bold=True)
    ws['B4'].font = Font(bold=True)
    wb.move_sheet('Summary', offset=-(len(wb.sheetnames) - 2))

    # ---------------------------------------------------------------- Estimates vs paid
    paid = lambda v: round(sum(x['amt'] for x in by_vendor[v]), 2)   # noqa: E731
    est = [
        ['SE Group due diligence', date(2025, 7, 15), 69_000, 75_900, paid('Sno.engineering'),
         '"We anticipate our professional fees will not exceed $69,000," plus expenses of "typically 10%." Res. '
         '2025-15 approved the agreement "substantially in the form of the proposal."',
         'Proposal pp. 4–5', DOC.format(122625)],
        ['Plante Moran transition planning', date(2026, 4, 21), 110_000, 110_000, paid('P&M Holding Group'),
         'Estimated $90,000 to $110,000 over four to five weeks. Its Eldora acquisition charges are on invoices '
         'dated Dec. 4, 2025, to April 17, 2026. Res. 2025-22 (Oct. 7, 2025) authorized the Town Manager to engage HR '
         'and finance consultants; it names no firm or estimate.', 'Engagement letter p. 3', DOC.format(128983)],
        ['Dickinson Wright water-rights due diligence', date(2025, 7, 15), None, None, paid('Dickinson Wright PLLC'),
         'Hourly rate only ($385); the engagement letter gives no estimate.',
         'Memo', DOC.format(122662)],
        ['Town Attorney, Widner Juran', date(2025, 7, 15), None, None, paid('Widner Juran LLP'),
         'A July 2025 rate memo set a $290 hourly "pass-through" rate from January 2026. No estimate of hours.',
         'Memo', DOC.format(122661)],
        ['Hogan Lovells (other legal counsel, Res. 2025-16)', date(2025, 8, 5), None, None,
         paid('Hogan Lovells US LLP'), 'The Town Manager "will keep the Board of Trustees informed about '
         'all necessary legal engagements." Hogan Lovells appears in no Board document searched except the warrant lists. A briefing in executive session would not appear in public records.',
         'Res. 2025-16', DOC.format(123040)],
        ['Stifel (financial advisor)', date(2026, 7, 7), 325_000, 325_000, 0,
         'Paid only if the bonds are sold.', 'Memo', DOC.format(130471)],
    ]
    sheet(wb, 'Estimates vs paid', ['Engagement', 'Approved or disclosed', 'Estimate (high end)',
                                    'Estimate incl. expenses', 'Paid', 'Notes',
                                    'Source', 'Source URL'], est, [34, 13, 14, 14, 14, 70, 18, 48],
          formats={'Approved or disclosed': DATE, 'Estimate (high end)': USD0, 'Estimate incl. expenses': USD0,
                   'Paid': USD},
          wrap={'Notes', 'Engagement'}, links={'Source': 'Source URL'})

    # ---------------------------------------------------------------- Commitments
    cm = [
        ["Purchase price for Eldora's assets", 120_000_000, '', 'Expired', 'No deposit required. The Town said Sept. '
         '30, 2026, that the agreement "has expired."', 'Ord. 882, APA §2.5', DOC.format(126803)],
        ['Series 2026 revenue bonds (up to 40 years, at most 9% interest)', 225_000_000, 'up to', 'Lapsed',
         'Had to be sold within 180 days (about July 19, 2026). None were sold.', 'Ord. 885 §5(c)',
         DOC.format(127180)],
        ['Expected maximum bond amount (intent to repay early costs from bonds)', 180_000_000, '', 'Declared', '',
         'Res. 2025-17', DOC.format(123060)],
        ['Transaction costs owed if the sale does not close', ESTIMATE, 'about', 'Town estimate', 'To be paid from '
         'CSAFE investments; reimbursable from bond proceeds if the deal closes.', 'Sept. 30, 2026, update',
         TOWN_UPDATE],
        ['Planned philanthropic fundraising for capital projects', 20_000_000, 'about', 'Planned', 'Jolly Jug, Corona '
         'lift, mountaintop lodge, summer operations.', 'Sept. 30, 2026, update', TOWN_UPDATE],
        ['Stifel financial-advisor fee', 325_000, '', 'If bonds sell', 'Paid from bond proceeds.',
         'July 7, 2026, memo', DOC.format(130471)],
        ['Emergency spending beyond the budget (Res. 2026-13)', None, 'no dollar limit', 'Lapsed', 'Authority ran 120 '
         'days from April 21, 2026, or until a supplemental budget.', 'Res. 2026-13', DOC.format(129046)],
        ['Mountain Recreation Enterprise in the 2025 and 2026 budgets', 0, '', 'Not budgeted', '',
         'April 21, 2026, memo', DOC.format(129045)],
    ]
    sheet(wb, 'Commitments', ['Item', 'Amount', 'Qualifier', 'Status', 'Notes', 'Source', 'Source URL'], cm,
          [52, 16, 14, 14, 60, 20, 48], formats={'Amount': USD0}, wrap={'Item', 'Notes'},
          links={'Source': 'Source URL'})

    # ---------------------------------------------------------------- CSAFE and funds
    funds = [
        ['Water and Sewer (utility funds)', 4_150_620, -1_118_188, 'Generally no', 'Utility rate revenue funds the utilities; the '
         "audit classifies $436,250 of the utilities' net position as restricted."],
        ['Downtown Development Authority', 1_695_802, 0, 'No', "A separate authority's money for downtown purposes."],
        ['Community Center', 964_480, -158_347, 'No', 'Restricted to the community center.'],
        ['Streets; Conservation Trust (parks)', 343_690, -179_698, 'No', 'Restricted to roads, parks, and trails.'],
        ['General Fund', 0, 1_456_233, 'Yes, if the Board appropriates it', 'Its cash was negative: the audit says its '
         '"negative cash balance of $757,482" is "reported as accounts payable." $147,000 of its fund balance is '
         'restricted as the TABOR emergency reserve.'],
    ]
    owed = 'Amount it is owed by (+) or owes (-) other funds'
    ws = sheet(wb, 'CSAFE and funds', ['Fund (Dec. 31, 2025)', 'Cash and investments', owed,
                                       'Can it pay Eldora acquisition costs?', 'Why'], funds, [34, 16, 22, 22, 80],
               formats={'Cash and investments': USD0, owed: USD0}, wrap={'Why', 'Can it pay Eldora acquisition costs?'})
    ws.auto_filter.ref = f'A1:E{len(funds) + 1}'
    ws.append(['Total', f'=SUM(B2:B{len(funds) + 1})', f'=SUM(C2:C{len(funds) + 1})', '',
               "Cash total equals the audit's deposits and investments ($7,154,592); amounts owed and due net to zero."])
    for col in (2, 3):
        ws.cell(ws.max_row, col).number_format = USD0
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(['If the other funds paid the General Fund what they owed, and that covered its negative cash balance, '
               'the General Fund would hold about $698,751 of the pooled cash and investments ($1,456,233 - $757,482; '
               "this workbook's calculation, not an audit figure). The audit expects the amounts owed to clear after "
               "year end but doesn't say how the negative balance would be settled. Eldora acquisition costs paid "
               'since year end are on the Payments sheet.'])
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.append([])
    ws.append(['Key audited figures', 'Amount', 'Date', 'Source'])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for label, v, when in [
        ('CSAFE investment balance', 2_700_831, 'Dec. 31, 2024'),
        ('CSAFE investment balance', 2_821_014, 'Dec. 31, 2025'),
        ('General Fund fund balance', 1_030_078, 'Dec. 31, 2025'),
        ('General Fund unassigned fund balance', 872_173, 'Dec. 31, 2025'),
        ("General Fund negative cash balance (reported as accounts payable)", -757_482, 'Dec. 31, 2025'),
        ('Owed to the General Fund by other funds', 1_456_233, 'Dec. 31, 2025'),
        ('General Fund TABOR emergency reserve', 147_000, 'Dec. 31, 2025'),
        ('General Fund "other receivables"', 6_431, 'Dec. 31, 2024'),
        ('General Fund "other receivables" (not itemized in the audit)', 520_055, 'Dec. 31, 2025'),
        ('General Fund balance projected in the adopted 2026 budget', 583_730, 'Dec. 31, 2026'),
    ]:
        ws.append([label, v, when, '2025 audit' if 'projected' not in label else '2026 budget summary'])
        ws.cell(ws.max_row, 2).number_format = USD0
        ws.cell(ws.max_row, 4).hyperlink = AUDIT_2025 if 'projected' not in label else DOC.format(126178)

    # ---------------------------------------------------------------- For scale
    sc = [['Eldora, if the sale fails (Town estimate)', ESTIMATE, 0, '=B2-C2', 'no grants', 'Sept. 30, 2026, update',
           TOWN_UPDATE],
          ['Eldora, paid so far', paid_total, 0, '=B3-C3', 'no grants', 'Summary sheet', '']]
    for label, total, grant, doc, note in page.SCALE_PROJECTS:
        n = len(sc) + 2
        note = re.sub(r'\s*\[\[corrected:[^\]]*\]\]', '', note)   # the page's correction marker
        sc.append([label, total, grant, f'=B{n}-C{n}', note, 'Board packet', DOC.format(doc)])
    ws = sheet(wb, 'For scale', ['Project', 'Total', 'Grant share', 'Town share', 'Note', 'Source', 'Source URL'], sc,
               [40, 15, 15, 15, 26, 22, 48], formats={k: USD0 for k in ('Total', 'Grant share', 'Town share')},
               links={'Source': 'Source URL'})
    ws.auto_filter.ref = f'A1:G{len(sc) + 1}'
    ws.append([])
    ws.append(['Town budgets', 'Amount', '', '', 'Note', 'Source', 'Source URL'])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for label, v, note, doc in [
        ('All funds, adopted 2026 budget', 9_400_000, 'about; includes ~$2.25M in grant-funded projects', 126174),
        ('General Fund operating spending, 2026 budget', 3_236_824, '', 126178),
        ('Boulder County Sheriff contract (largest General Fund cost)', 1_050_000, 'about, per year', 116150),
        ('General Fund capital spending', 200_000, '$160,000–$200,000 a year', 126178),
    ]:
        ws.append([label, v, None, None, note, 'Board packet', DOC.format(doc)])
        ws.cell(ws.max_row, 2).number_format = USD0
        ws.cell(ws.max_row, 6).hyperlink = DOC.format(doc)

    # ---------------------------------------------------------------- Board votes
    votes = list(csv.DictReader(open(os.path.join(ROOT, 'analysis', 'eldora_votes.csv'))))
    cols = ['date', 'category', 'description', 'outcome', 'vote', 'mover', 'seconder', 'absent', 'minutes_page',
            'source', 'motion_text']
    vr = [[d(v['date'])] + [v[c] for c in cols[1:]] for v in votes]
    sheet(wb, 'Board votes', ['Date', 'Category', 'Description', 'Outcome', 'Vote', 'Mover', 'Seconder', 'Absent',
                              'Minutes page', 'Source URL', 'Motion text'], vr,
          [12, 18, 60, 10, 40, 12, 12, 14, 9, 48, 80], formats={'Date': DATE},
          wrap={'Description', 'Vote', 'Motion text'})

    path = os.path.join(OUT, 'eldora-record.xlsx')
    tmp = path.replace('.xlsx', '.new.xlsx')
    wb.save(tmp)
    if os.path.exists(path) and contents(path) == contents(tmp):
        os.remove(tmp)
        print(f'{path}: unchanged')
    else:
        os.replace(tmp, path)
        print(f'{path}: NEW VERSION. Upload it to the Drive copy (Manage versions > Upload new version).')

    # independent check of the formulas' expected results
    listed = sum(r[5] for r in pay_rows if r[6] == 'Yes')
    off = sum(g[2] for g in gaps)
    print(f'  {len(wb.sheetnames)} sheets, {n_pay} payment rows')
    print(f'  listed {listed:,.2f} + off-list {off:,.2f} = {listed + off:,.2f}; last balance {paid_total:,.2f}; '
          f'unattributed {unattributed:,.2f}; not yet paid {ESTIMATE - paid_total:,.2f}')
    assert abs(listed + off - paid_total) < 0.01


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    main()
